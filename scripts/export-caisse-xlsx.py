import json, io, sys, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

S = sys.argv[1]
d = json.load(io.open(S + '/export-caisse.json', encoding='utf-8'))
recettes, ets, opts, plats, tags = d['recettes'], d['ets'], d['opts'], d['plats'], d['tags']
fab = {f['id']: f['name'] for f in (d.get('fab') or {}).get('data', [])}
POSTES = (d.get('fab') or {}).get('data', [])
PAIEMENTS = (d.get('paiements') or {}).get('data', []) or (d.get('paiements') or {}).get('transaction_methods', [])
TABLES, CLIENTS, EMPLOYES = d.get('tables') or [], d.get('clients') or [], d.get('employes') or []
NOM_ET = {e['id']: e['nom'] for e in ets}
PAR_REMOTE = {str(p.get('remote_id') or ''): p for p in plats}
NOM_TAG = {t['id']: t['name'] for t in tags}
ORDRE_TAG = {t['name']: (t.get('o') or 0) for t in tags}
NOM_GRP = {g['id']: g['name'] for g in opts}

ARIAL = 'Arial'
TITRE = Font(name=ARIAL, size=11, bold=True, color='FFFFFF')
FOND  = PatternFill('solid', fgColor='253328')          # vert CasaTasia
GRAS  = Font(name=ARIAL, size=11, bold=True)
NORM  = Font(name=ARIAL, size=10)
ALERTE= Font(name=ARIAL, size=10, bold=True, color='B00020')
BAS   = Border(bottom=Side('thin', color='D9D9D9'))

wb = Workbook()

def feuille(titre, entetes, lignes, largeurs):
    ws = wb.create_sheet(titre)
    ws.append(entetes)
    for c in ws[1]:
        c.font, c.fill = TITRE, FOND
        c.alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 30
    for l in lignes:
        ws.append(l)
    for i, w in enumerate(largeurs, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = NORM
            c.border = BAS
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    return ws

# ───────────────────────── 1. LIRE-MOI ─────────────────────────
ws = wb.active; ws.title = 'LIRE-MOI'
ws.column_dimensions['A'].width = 4
ws.column_dimensions['B'].width = 118
def L(txt='', style=None, h=None):
    ws.append(['', txt])
    c = ws.cell(ws.max_row, 2)
    c.font = style or NORM
    c.alignment = Alignment(wrap_text=True, vertical='top')
    if h: ws.row_dimensions[ws.max_row].height = h

L('CASATASIA — export du catalogue de caisse', Font(name=ARIAL, size=16, bold=True, color='253328'))
L('Parking des Ferrages · 83136 Sainte-Anastasie-sur-Issole · casatasia.fr')
L(f"Export du {datetime.date.today().strftime('%d/%m/%Y')}")
L()
L('À LIRE AVANT L’IMPORT — cinq points qui ne se devinent pas', GRAS)
L()
L('1. CHAQUE PRODUIT A DEUX PRIX, et ce ne sont pas des arrondis.', ALERTE)
L('   « Prix TTC sur place » et « Prix TTC à emporter » sont deux colonnes distinctes. 21 produits ont')
L('   un écart réel : un Coca 33 cl est à 2,00 € au comptoir et 3,20 € servi en salle (verre consigné).')
L('   N’en importer qu’un seul reviendrait à vendre la salle au tarif du comptoir.', None, 30)
L()
L('2. CHAQUE PRODUIT A DEUX TAUX DE TVA, imposés par la loi française.', ALERTE)
L('   À emporter : 5,5 % (boulangerie) ou 10 % (snacking, pizzas, boissons). Sur place : 10 %.')
L('   L’alcool reste à 20 % dans les deux modes. Le non alimentaire (famille Tabac) garde son taux.')
L('   Un taux trop BAS est le seul qui ne se rattrape pas : on ne réclame pas après coup au client.', None, 30)
L()
L('3. DEUX GROUPES D’OPTIONS SONT OBLIGATOIRES, pas facultatifs.', ALERTE)
L('   « Cuisson » (entrecôte) et « Cuisson du burger » ont min = 1 : la caisse doit EXIGER le choix.')
L('   Et ce sont deux groupes distincts à dessein — « bleu » ne se sert pas sur un steak haché.')
L('   Le menu enfant porte le même steak haché et ne reçoit AUCUN groupe : sa cuisson à cœur')
L('   n’est pas une option.', None, 30)
L()
L('4. LES PRODUITS « ÉTEINTS » NE DOIVENT PAS ÊTRE MIS EN VENTE.', ALERTE)
L('   L’onglet Produits les liste avec Actif = NON. Ils sont conservés exprès : un ticket tardif doit')
L('   pouvoir s’y rattacher, et certains reviennent à la carte (les plats de brasserie tournent chaque')
L('   semaine sur une ardoise). Ne pas les supprimer, ne pas les activer.', None, 30)
L()
L('5. GARDEZ NOTRE CODE PRODUIT, c’est la clé de rapprochement.', ALERTE)
L('   La colonne « Code CasaTasia » est notre identifiant interne et stable. S’il est conservé dans un')
L('   champ libre de votre caisse, la correspondance entre les deux systèmes reste exacte même si un')
L('   nom ou un prix change plus tard. Sans lui, tout rapprochement se fait par le libellé — et un')
L('   libellé qui change coupe la série statistique en deux, sans erreur visible.', None, 42)
L()
L('CE QUE CET EXPORT NE CONTIENT PAS, volontairement', GRAS)
L('   Aucun prix d’achat, aucun fournisseur, aucune référence fournisseur : ce sont des conditions')
L('   négociées, et une caisse n’en a pas besoin pour encaisser.', None, 30)
L()
L('CONTENU DU CLASSEUR', GRAS)

# ───────────────────────── 2. PRODUITS ─────────────────────────
ALL14 = ['gluten','crustaces','oeufs','poissons','arachides','soja','lait','fruits_a_coque',
         'celeri','moutarde','sesame','sulfites','lupin','mollusques']
LIB_ALL = {'gluten':'Gluten','crustaces':'Crustacés','oeufs':'Œufs','poissons':'Poissons',
           'arachides':'Arachides','soja':'Soja','lait':'Lait','fruits_a_coque':'Fruits à coque',
           'celeri':'Céleri','moutarde':'Moutarde','sesame':'Sésame','sulfites':'Sulfites',
           'lupin':'Lupin','mollusques':'Mollusques'}

lignes, sans_caisse = [], []
for r in recettes:
    z = PAR_REMOTE.get(str(r['id']))
    if z is None:
        sans_caisse.append(r['nom'])
    # ⚠️ prix et taux LUS SUR LA CAISSE : ils y sont déjà contrôlés (209/209
    # conformes). Les recalculer rouvrirait le risque d'un arrondi divergent.
    if z:
        ttc_salle = round((z.get('price') or 0) / 100, 2)
        ttc_emp   = round((z.get('price_togo') or z.get('price') or 0) / 100, 2)
        tva_salle = (z.get('tax') or 0) / 100
        tva_emp   = (z.get('tax_takeaway') or z.get('tax') or 0) / 100
        poste     = fab.get(z.get('id_fabrication_place')) or ''
        groupes   = ' + '.join(NOM_GRP.get(g, str(g)) for g in (z.get('options') or []))
        fams      = ' + '.join(NOM_TAG.get(t if not isinstance(t, dict) else t.get('id'), '?')
                               for t in (z.get('tags') or []))
    else:
        base = round(r['prix_vente_ht'] * (1 + r['tva'] / 100), 2)
        ttc_emp   = base
        ttc_salle = round(float(r['prix_sur_place_ttc']), 2) if r.get('prix_sur_place_ttc') else base
        tva_emp   = float(r['tva'])
        tva_salle = 20.0 if r['contient_alcool'] else (float(r['tva']) if r['categorie'] == 'Tabac' else 10.0)
        poste, groupes, fams = '', '', r['categorie']
    alg = r.get('allergenes_complementaires') or []
    lignes.append([
        r['id'],                                   # Code CasaTasia
        r.get('nom_caisse') or r['nom'],           # Libellé ticket
        r['nom'],                                  # Nom de vitrine
        fams or r['categorie'],                    # Famille(s) caisse
        NOM_ET.get(r.get('etablissement_id'), ''), # Point de vente
        ttc_salle, ttc_emp,
        tva_salle, tva_emp,
        'OUI' if r['contient_alcool'] else '',
        'OUI' if r['actif'] else 'NON',
        'OUI' if r.get('vendable_online') else '',
        poste,
        groupes,
        ', '.join(LIB_ALL.get(a, a) for a in alg),
        'signé' if r.get('allergenes_valides_le') else ('déclaré, non signé' if alg else 'non vérifié'),
        r.get('image_url') or '',
        (r.get('description') or '')[:200],
    ])

# ⚠⚠ DEUX ONGLETS, PAS UN. 290 lignes dont 81 éteintes, c'est du bruit pour
# qui doit importer : l'intégrateur prend « Produits » et ne peut pas activer par
# erreur une fiche morte. Et les éteintes ne se valent pas — celles de la
# restauration REVIENDRONT (la brasserie tourne chaque semaine sur une ardoise),
# les autres sont des retraits décidés.
RESTO = {'Burger','Plat','Planche','Grande salade','Menu','Pizzeria'}
ENTETES = ['Code CasaTasia','Libellé ticket','Nom de vitrine','Famille(s)','Point de vente',
   'Prix TTC sur place','Prix TTC à emporter','TVA sur place %','TVA emporter %',
   'Alcool','Actif','Vente en ligne','Poste de production','Groupes d’options',
   'Allergènes déclarés','État de la déclaration','URL photo','Description']
LARG = [38,34,34,26,15,12,12,10,10,8,7,9,16,34,30,20,62,46]

actifs  = [l for l in lignes if l[10] == 'OUI']
eteints = [l for l in lignes if l[10] == 'NON']

# ⚠⚠ LES EN-TÊTES DOIVENT PERDRE LA MÊME COLONNE QUE LES LIGNES. Première
# version : `ENTETES` complet (18) pour des lignes à 17 colonnes — l'en-tête
# « Actif » restait et tout glissait d'un cran à partir de là, si bien qu'un
# groupe d'options s'affichait sous « Poste de production ». Aucune erreur :
# un décalage de colonne rend des valeurs PLAUSIBLES à la mauvaise place, et
# c'est l'intégrateur qui l'aurait importé tel quel.
feuille('Produits', ENTETES[:10] + ENTETES[11:], [l[:10] + l[11:] for l in actifs],
        LARG[:10] + LARG[11:])

# la colonne « Actif » n'a pas de sens dans un onglet qui ne contient que des
# actifs : elle est retirée ci-dessus, et remplacée ici par la RAISON du retrait.
CAT = {r['id']: r['categorie'] for r in recettes}
et_l = []
for l in eteints:
    fam_base = CAT.get(l[0], '')
    et_l.append(l[:10] + l[11:] + [
        'REVIENDRA — tourne à l’ardoise de la semaine' if fam_base in RESTO
        else 'retiré de la carte — ne pas activer'])
feuille('Éteints (ne pas vendre)', ENTETES[:10] + ENTETES[11:] + ['Pourquoi'],
        et_l, LARG[:10] + LARG[11:] + [46])

# ───────────────────────── 3. FAMILLES ─────────────────────────
par_fam = {}
for r in recettes:
    if not r['actif']: continue
    par_fam.setdefault(r['categorie'], []).append(r)
fl = []
for t in sorted(tags, key=lambda t: (t.get('o') or 0)):
    n = t['name']
    l = par_fam.get(n, [])
    taux = sorted({float(x['tva']) for x in l})
    fl.append([t.get('o') or 0, n, len(l),
               ' / '.join(f'{x:g}' for x in taux) if taux else '',
               'rayon VIDE — aucun produit actif' if not l else ''])
feuille('Familles', ["Ordre d’affichage",'Famille','Produits actifs','TVA emporter présentes %','Remarque'],
        fl, [16,30,15,24,40])

# ───────────────────────── 4. OPTIONS ─────────────────────────
ol = []
for g in opts:
    for v in (g.get('values') or []):
        ol.append([g['name'], g['id'],
                   'OBLIGATOIRE' if (g.get('min_choices') or 0) >= 1 else 'facultatif',
                   g.get('min_choices'), g.get('max_choices'),
                   'oui' if g.get('pick_same_value') else 'non',
                   "s’affiche à chaque vente" if not g.get('donotpresent') else 'à la demande du serveur',
                   v['name'], round((v.get('price') or 0) / 100, 2)])
feuille('Options',
  ['Groupe','Id groupe','Nature','Min','Max','Même valeur plusieurs fois',
   'Affichage','Valeur','Prix TTC du supplément'],
  ol, [26,10,14,6,6,24,26,30,20])

# ───────────────────── 5. OPTIONS PAR PRODUIT ─────────────────
rl = []
for r in recettes:
    z = PAR_REMOTE.get(str(r['id']))
    for gid in ((z or {}).get('options') or []):
        rl.append([r.get('nom_caisse') or r['nom'], r['id'], NOM_GRP.get(gid, str(gid)), gid,
                   'OBLIGATOIRE' if any((g.get('min_choices') or 0) >= 1 for g in opts if g['id'] == gid) else 'facultatif'])
feuille('Options par produit', ['Libellé ticket','Code CasaTasia','Groupe','Id groupe','Nature'],
        rl, [34,38,26,10,14])

# ───────────────────────── 6. ALLERGÈNES ─────────────────────
al = []
for a in ALL14:
    pr = [r for r in recettes if r['actif'] and a in (r.get('allergenes_complementaires') or [])]
    al.append([LIB_ALL[a], a, len(pr), ', '.join(x['nom'] for x in pr[:40]) + (' …' if len(pr) > 40 else '')])
feuille('Allergènes', ['Allergène (14 réglementaires)','Clé','Produits actifs concernés','Lesquels'],
        al, [30,16,22,110])

# ───────────────── 7. MODES DE PAIEMENT ──────────────────────
# Une caisse ne peut rien solder sans eux, et il n'en existait AUCUN avant
# septembre — pas même les espèces. Le type « account » est celui qui permet la
# note EN COMPTE : la vente est faite, l'argent n'est pas rentré. Il doit rester
# DISTINCT des espèces et de la CB, sinon le comptage du tiroir crie un écart
# chaque soir.
TYPE_FR = {'cash': 'espèces (rendu monnaie, tiroir, comptage)',
           'ingenico': 'TPE Ingenico — le montant part de la caisse vers le terminal',
           'tr': 'titre-restaurant', 'account': 'EN COMPTE — ardoise client, rien n’est encaissé'}
pl = [[m.get('order') or 0, m['name'], m.get('type') or '(aucun)', TYPE_FR.get(m.get('type') or '', ''),
       'oui' if m.get('open_drawer') else 'non',
       {'change': 'rendu monnaie', 'allow': 'autorisé', 'disallow': 'interdit'}.get(m.get('overflow') or '', m.get('overflow') or ''),
       'oui' if m.get('allow_negative_transacion') else 'non',
       'DÉSACTIVÉ' if m.get('disable') else 'actif',
       m.get('accounting_code') or '⚠ à définir avec le comptable']
      for m in sorted(PAIEMENTS, key=lambda m: m.get('order') or 0)]
feuille('Modes de paiement', ['Ordre','Mode','Type technique','Ce que le type implique','Ouvre le tiroir',
                              'Surplus','Montant négatif','État','Code comptable'],
        pl, [8,22,16,52,15,15,16,12,30])

# ───────────────── 8. TABLES ET ZONES ────────────────────────
tl = [[t.get('zone') or '—', t.get('numero'), t.get('capacite') or 0] for t in TABLES]
zones = {}
for t in TABLES:
    z = t.get('zone') or '—'
    zones[z] = (zones.get(z, (0, 0))[0] + 1, zones.get(z, (0, 0))[1] + (t.get('capacite') or 0))
ws = feuille('Tables et zones', ['Zone','Table','Couverts'], tl, [18,12,12])
ws.append([]); ws.append(['TOTAL', f'{len(TABLES)} tables', sum(t.get('capacite') or 0 for t in TABLES)])
for c in ws[ws.max_row]: c.font = GRAS
for z, (n, cv) in sorted(zones.items()):
    ws.append([z, f'{n} tables', cv])
    for c in ws[ws.max_row]: c.font = NORM

# ───────────────── 9. COMPTES CLIENTS ────────────────────────
# ⚠️ NI allergies (donnée de SANTÉ) NI consentement marketing : un consentement
# ne se transfère pas d'un prestataire à l'autre, et il se recueille — il ne
# s'hérite pas d'une valeur par défaut. Le total est en HT (base de la fidélité).
cl = [[f"{(c.get('prenom') or '').strip()} {(c.get('nom') or '').strip()}".strip(),
       c.get('telephone') or '', c.get('email') or '',
       c.get('caisse_externe_id') or '', c.get('total_depense') or 0, c.get('nb_visites') or 0,
       c.get('niveau_fidelite') or '', c.get('points_fidelite') or 0]
      for c in CLIENTS]
feuille('Comptes clients', ['Nom','Téléphone','Email','Id caisse actuelle','Total dépensé € HT',
                            'Visites','Niveau fidélité','Points'],
        cl, [30,16,32,18,18,10,16,10])
ws = wb['Comptes clients']
ws.append([])
ws.append(['⚠ Données personnelles : cet onglet peut être SUPPRIMÉ avant envoi. Le seul compte utile à la '
           'caisse est le compte PRO, qui porte une note en compte.'])
ws.cell(ws.max_row, 1).font = ALERTE

# ───────────────── 10. POSTES ET IMPRIMANTES ─────────────────
# Les trois Epson TM-T88VI sont relevées à la main (ticket d'état de chaque
# imprimante) : elles ne sont dans aucune base. ⚠️ Les IP sont en DHCP et
# doivent être RÉSERVÉES par adresse MAC dans la box — l'iPad est déjà passé de
# .69 à .199, et la même dérive sur une imprimante enverrait les bons de la
# cuisine chez le pizzaiolo, sans le moindre signalement.
IMPR = [['Comptoir', 'imprimante 80 mm', '192.168.1.196', 'MAC 44:D2:44:C4:61:2E · câble Ethernet sur la box', 'Epson TM-T88VI'],
        ['Cuisine',  'imprimante 80 mm', '192.168.1.47',  'nom réseau EPSON199BBE · via le répéteur TP-Link RE505X', 'Epson TM-T88VI'],
        ['Pizza',    'imprimante 80 mm', 'à brancher',    'il manque un switch 5 ports', 'Epson TM-T88VI']]
ws = feuille('Postes et imprimantes', ['Nom','Nature','Adresse IP','MAC / raccordement','Matériel'],
             [[p['name'], 'poste de production', '', '', ''] for p in POSTES] + IMPR,
             [20,24,18,32,22])
ws.append([])
ws.append(['', 'Réseau : box 192.168.1.254, SSID CASATASIA. Les IP sont en DHCP — à réserver par MAC.'])
ws.cell(ws.max_row, 2).font = ALERTE

# ── le sommaire du Lire-moi, écrit après coup pour porter les vrais comptes ──
ws = wb['LIRE-MOI']
for t, n in [('Produits', f'{len(actifs)} produits ACTIFS — c’est l’onglet à importer'),
             ('Éteints (ne pas vendre)', f'{len(eteints)} fiches, avec la raison du retrait'),
             ('Familles', f'{len(fl)} rayons, dans leur ordre d’affichage'),
             ('Options', f'{len(ol)} valeurs réparties en {len(opts)} groupes'),
             ('Options par produit', f'{len(rl)} rattachements'),
             ('Allergènes', 'les 14 réglementaires et les produits concernés'),
             ('Modes de paiement', f'{len(pl)} modes, dont la note EN COMPTE'),
             ('Tables et zones', f'{len(TABLES)} tables, {sum(t.get("capacite") or 0 for t in TABLES)} couverts'),
             ('Comptes clients', f'{len(cl)} comptes (sans donnée de santé ni consentement marketing)'),
             ('Postes et imprimantes', f'{len(POSTES)} postes de production et 3 imprimantes 80 mm')]:
    ws.append(['', f'   · {t} — {n}'])
    ws.cell(ws.max_row, 2).font = NORM
if sans_caisse:
    ws.append(['', ''])
    ws.append(['', f'⚠ {len(sans_caisse)} produit(s) sans contrepartie dans la caisse actuelle : ' + ', '.join(sans_caisse[:10])])
    ws.cell(ws.max_row, 2).font = ALERTE

out = 'data/export-caisse-casatasia-' + datetime.date.today().strftime('%Y-%m-%d') + '.xlsx'
wb.save(out)
print('✓ écrit :', out)
print(f'  Produits {len(lignes)} · Familles {len(fl)} · Options {len(ol)} · Rattachements {len(rl)}')
print(f'  sans contrepartie caisse : {len(sans_caisse)}')
print(f'  Paiements {len(pl)} · Tables {len(TABLES)} · Clients {len(cl)} · Postes {len(POSTES)}')
