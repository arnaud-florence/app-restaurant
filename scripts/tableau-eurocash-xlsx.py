# Le tableau Euro-Cash, en classeur à envoyer par mail.
#
#   node scripts/selection-eurocash.mjs   # produit le CSV
#   python3 scripts/tableau-eurocash-xlsx.py
#
# Le CSV est la source ; ce script ne fait que le mettre en forme. Recalculer
# les prix ici en dupliquerait la logique, et les deux finiraient par diverger.
#
# ⚠️ Les colonnes « VOTRE PRIX » et « VOTRE REMISE » sont laissées en JAUNE et
# vides : un fournisseur doit voir au premier coup d'œil ce qu'on attend de
# lui. Un classeur où tout se ressemble revient rempli à moitié.
#
# ⚠️ Le classeur reste dans data/, gitignoré : il porte les prix que nous
# payons chez Gineys, Promocash et France Boissons.

import csv, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC = 'data/tarif-eurocash-2026-09-23.csv'
OUT = 'data/Demande-de-tarif-CASATASIA-Euro-Cash.xlsx'

rows = list(csv.reader(open(SRC, encoding='utf-8-sig'), delimiter=';'))
# Le fichier porte DEUX tableaux séparés par une ligne vide : les références,
# puis les besoins sans équivalent. Les séparer ici plutôt que de les empiler
# sur une même feuille — un commercial remplit une colonne, il ne cherche pas
# où le premier tableau s'arrête.
coupe = next(i for i, r in enumerate(rows) if not any(c.strip() for c in r))
refs, entete_refs = rows[1:coupe], rows[0]
bloc2 = rows[coupe+1:]
titre2, entete2, lignes2 = bloc2[0][0], bloc2[1], bloc2[2:]

ARIAL   = 'Arial'
BLEU    = PatternFill('solid', fgColor='1F3864')
JAUNE   = PatternFill('solid', fgColor='FFF2CC')
VERT    = PatternFill('solid', fgColor='E2EFDA')
GRIS    = PatternFill('solid', fgColor='F2F2F2')
FIN     = Side(style='thin', color='BFBFBF')
BORDURE = Border(left=FIN, right=FIN, top=FIN, bottom=FIN)

wb = openpyxl.Workbook()

# ── Feuille 1 : les références ───────────────────────────────────
ws = wb.active
ws.title = 'Demande de tarif'

ws['A1'] = 'CASATASIA — Demande de tarif Euro-Cash'
ws['A1'].font = Font(name=ARIAL, size=15, bold=True, color='1F3864')
ws['A2'] = ('Boulangerie, bar, revente tabac et relais colis — Sainte-Anastasie-sur-Issole (83136). '
            'Catalogue général été 2026. Merci de compléter les deux dernières colonnes.')
ws['A2'].font = Font(name=ARIAL, size=10, italic=True, color='595959')
ws['A3'] = ('« Notre prix actuel » et « Prix colis à battre » indiquent ce que nous payons aujourd’hui ailleurs, '
            'ramené à votre colisage. Une case vide = nous n’avons pas encore de fournisseur sur cette référence.')
ws['A3'].font = Font(name=ARIAL, size=10, color='C00000')
for r in (1, 2, 3):
    ws.cell(row=r, column=1).alignment = Alignment(vertical='center')

LIGNE_ENTETE = 5
for j, t in enumerate(entete_refs, start=1):
    c = ws.cell(row=LIGNE_ENTETE, column=j, value=t)
    c.font = Font(name=ARIAL, size=10, bold=True, color='FFFFFF')
    c.fill = BLEU
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    c.border = BORDURE
ws.row_dimensions[LIGNE_ENTETE].height = 34

NUM = {7, 9, 10, 11}   # colonnes de montants
for i, r in enumerate(refs, start=LIGNE_ENTETE + 1):
    prioritaire = r[1].strip() == 'PRIORITAIRE'
    for j, v in enumerate(r, start=1):
        val = v
        if j in NUM and v.strip():
            try: val = float(v.replace(',', '.'))
            except ValueError: pass
        c = ws.cell(row=i, column=j, value=val)
        c.font = Font(name=ARIAL, size=10, bold=(prioritaire and j == 2))
        c.border = BORDURE
        if j in NUM: c.number_format = '#,##0.0000 €' if j == 7 else '#,##0.00 €'
        if j in (10, 11): c.fill = JAUNE            # à remplir par Euro-Cash
        elif prioritaire: c.fill = VERT
        elif i % 2 == 0: c.fill = GRIS
    ws.cell(row=i, column=3).alignment = Alignment(horizontal='left')

for col, w in zip('ABCDEFGHIJK', (42, 12, 9, 46, 14, 8, 17, 34, 16, 14, 14)):
    ws.column_dimensions[col].width = w
ws.freeze_panes = f'A{LIGNE_ENTETE + 1}'
ws.auto_filter.ref = f'A{LIGNE_ENTETE}:K{LIGNE_ENTETE + len(refs)}'

# ── Feuille 2 : les besoins sans équivalent ──────────────────────
w2 = wb.create_sheet('Avez-vous l’équivalent')
w2['A1'] = titre2
w2['A1'].font = Font(name=ARIAL, size=13, bold=True, color='1F3864')
w2['A2'] = ('Ces références sont achetées ailleurs aujourd’hui. Nous n’avons pas trouvé '
            'le format équivalent à votre catalogue — si vous l’avez, l’affaire est à prendre.')
w2['A2'].font = Font(name=ARIAL, size=10, italic=True, color='595959')

E2 = ['Désignation', 'Par', 'Notre prix HT du colis', 'Remarque', 'VOTRE PRIX HT', 'VOTRE REMISE']
for j, t in enumerate(E2, start=1):
    c = w2.cell(row=4, column=j, value=t)
    c.font = Font(name=ARIAL, size=10, bold=True, color='FFFFFF')
    c.fill = BLEU
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    c.border = BORDURE
w2.row_dimensions[4].height = 30

for i, r in enumerate(lignes2, start=5):
    vals = [r[0], r[1], r[2], r[3], '', '']
    for j, v in enumerate(vals, start=1):
        val = v
        if j in (2, 3) and str(v).strip():
            try: val = float(str(v).replace(',', '.'))
            except ValueError: pass
        c = w2.cell(row=i, column=j, value=val)
        c.font = Font(name=ARIAL, size=10)
        c.border = BORDURE
        if j == 3: c.number_format = '#,##0.000 €'
        if j == 2: c.number_format = '#,##0'
        if j in (5, 6): c.fill = JAUNE
    w2.cell(row=i, column=4).alignment = Alignment(wrap_text=True, vertical='top')

for col, w in zip('ABCDEF', (46, 8, 20, 58, 15, 15)):
    w2.column_dimensions[col].width = w
w2.freeze_panes = 'A5'

wb.save(OUT)
print('écrit :', OUT)
print('feuille 1 :', len(refs), 'références |  feuille 2 :', len(lignes2), 'lignes')
