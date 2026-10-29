#!/usr/bin/env python3
"""Classeur Excel d'un relevé client, à partir du CSV + .meta.json produits
par `releve-client.mjs`. Même patron que `export-caisse.mjs` → `-xlsx.py` :
node sort les données, python met en forme.

    node scripts/releve-client.mjs --nom="…" --mois=2026-10 --csv
    python3 scripts/releve-client-xlsx.py data/releve-….csv

⚠️ LES MONTANTS SONT ÉCRITS COMME NOMBRES, jamais comme texte : un total
qu'Excel ne sait pas sommer est inutile sur le document qui justifie une
facture — c'est le même défaut que la TVA en « 5.5 ».

⚠️ LE CLASSEUR PORTE LA BORNE DU FLUX ET LES MANQUES DE LA FICHE. Un relevé
envoyé sans sa date de dernier ticket se lit comme un relevé à jour.
"""
import csv, json, sys, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

src = sys.argv[1]
meta = json.load(open(src.replace('.csv', '.meta.json'), encoding='utf-8'))
rows = list(csv.reader(open(src, encoding='utf-8-sig'), delimiter=';'))
head, body = rows[0], rows[1:]

def nb(t):
    t = (t or '').strip()
    return float(t.replace(',', '.')) if t else None

GRAS = Font(bold=True)
TITRE = Font(bold=True, size=14)
F_JOUR = PatternFill('solid', fgColor='EEF2FF')
F_TETE = PatternFill('solid', fgColor='253328')          # vert de marque
BAS = Border(bottom=Side(style='thin', color='CCCCCC'))

wb = Workbook()

# ───────── onglet 1 : le relevé, jour par jour ─────────
ws = wb.active
ws.title = 'Releve jour par jour'
ws['A1'] = f"RELEVÉ — {meta['client']}"
ws['A1'].font = TITRE
ws['A2'] = f"Période : {meta['periode']}   ·   {meta['tickets']} tickets sur {meta['journees']} journées"
if meta.get('borne'):
    ws['A3'] = f"⚠ Dernier ticket de caisse entré dans l'outil : {meta['borne']} — ce relevé est complet jusqu'à cette date, et muet au-delà."
    ws['A3'].font = Font(bold=True, color='B91C1C')
lig = 4
for m in meta.get('manques', []):
    ws.cell(lig, 1, f"⚠ {m}").font = Font(color='B45309')
    lig += 1
if meta.get('trous'):
    ws.cell(lig, 1, f"⚠ {len(meta['trous'])} jour(s) sans aucune ligne : " + ', '.join(meta['trous'])).font = Font(bold=True, color='B91C1C')
    lig += 1
else:
    ws.cell(lig, 1, f"✓ Aucun jour sans ligne : {meta['journees']} journées consécutives couvertes.")
    lig += 1

lig += 1
EN = ['Date', 'Jour', 'Ticket', 'Qté', 'Produit', 'PU TTC', 'Total TTC', 'TVA %', 'Règlement']
assert len(EN) == len(head), f"en-têtes : {len(EN)} contre {len(head)} colonnes au CSV"
for c, t in enumerate(EN, 1):
    cel = ws.cell(lig, c, t)
    cel.font = Font(bold=True, color='FFFFFF'); cel.fill = F_TETE
    cel.alignment = Alignment(horizontal='center')
lig += 1

for r in body:
    tot = r[2].startswith('TOTAL')
    # ⚠️ Autant de valeurs que d'en-têtes, sinon tout glisse d'un cran et les
    # valeurs restent PLAUSIBLES à la mauvaise place (vécu deux fois).
    vals = [r[0], r[1], r[2], nb(r[3]), r[4].strip('"'), nb(r[5]), nb(r[6]), nb(r[7]), r[8]]
    assert len(vals) == len(EN)
    for c, v in enumerate(vals, 1):
        cel = ws.cell(lig, c, v)
        if c in (6, 7):
            cel.number_format = '#,##0.00 €'
        if c == 8 and v is not None:
            cel.number_format = '0.0'
        if tot:
            cel.font = GRAS; cel.fill = F_JOUR
        cel.border = BAS
    lig += 1

for col, w in zip('ABCDEFGHI', (12, 12, 26, 6, 34, 11, 12, 8, 15)):
    ws.column_dimensions[col].width = w
ws.freeze_panes = ws.cell(lig - len(body), 1)

# ───────── onglet 2 : récapitulatif par produit ─────────
ws2 = wb.create_sheet('Par produit')
agg = {}
for r in body:
    if r[2].startswith('TOTAL'):
        continue
    k = (r[4].strip('"'), nb(r[5]), r[7])
    a = agg.setdefault(k, 0)
    agg[k] = a + (nb(r[3]) or 0)
ws2['A1'] = 'RÉCAPITULATIF PAR PRODUIT'; ws2['A1'].font = TITRE
ws2['A2'] = "Ce tableau sert à CONTRÔLER le relevé, pas à le remplacer : il ne dit pas quel jour la marchandise est partie."
for c, t in enumerate(['Qté', 'Produit', 'PU TTC', 'Total TTC', 'TVA %'], 1):
    cel = ws2.cell(4, c, t); cel.font = Font(bold=True, color='FFFFFF'); cel.fill = F_TETE
l = 5
for (nom, pu, tva), q in sorted(agg.items(), key=lambda x: -(x[1] * (x[0][1] or 0))):
    ws2.cell(l, 1, q); ws2.cell(l, 2, nom)
    ws2.cell(l, 3, pu).number_format = '#,##0.00 €'
    ws2.cell(l, 4, q * (pu or 0)).number_format = '#,##0.00 €'
    ws2.cell(l, 5, nb(tva)).number_format = '0.0'
    l += 1
ws2.cell(l + 1, 2, 'TOTAL').font = GRAS
ws2.cell(l + 1, 4, meta['ttc']).number_format = '#,##0.00 €'
ws2.cell(l + 1, 4).font = GRAS
ws2.cell(l + 2, 2, 'dont HT').font = GRAS
ws2.cell(l + 2, 4, meta['ht']).number_format = '#,##0.00 €'
for col, w in zip('ABCDE', (8, 36, 12, 13, 8)):
    ws2.column_dimensions[col].width = w

out = sys.argv[2] if len(sys.argv) > 2 else src.replace('.csv', '.xlsx')
wb.save(out)

# ⚠️ CONTRÔLE D'APRÈS ÉCRITURE : on relit le fichier produit. Un décalage de
# colonne ne lève aucune erreur, il met des valeurs plausibles au mauvais
# endroit — et c'est le destinataire qui le découvre.
from openpyxl import load_workbook
v = load_workbook(out)
w = v['Releve jour par jour']
tete = [c.value for c in next(w.iter_rows(min_row=1, max_row=w.max_row))]
ligne_en = next(i for i, r in enumerate(w.iter_rows(values_only=True), 1) if r[0] == 'Date')
ens = list(w.iter_rows(min_row=ligne_en, max_row=ligne_en, values_only=True))[0][:9]
assert list(ens) == EN, f"en-têtes décalés : {ens}"
somme = sum(r[6] for r in w.iter_rows(min_row=ligne_en + 1, values_only=True)
            if r[2] and not str(r[2]).startswith('TOTAL') and r[6])
assert abs(somme - meta['ttc']) < 0.01, f"somme des lignes {somme:.2f} ≠ total {meta['ttc']:.2f}"
for n in wb.sheetnames:
    assert len(n) <= 31, n   # au-delà, certaines applications refusent le fichier
print(f"✓ {out}")
print(f"  {len(body)} lignes · somme des lignes {somme:.2f} € = total {meta['ttc']:.2f} € ✓")
print(f"  onglets : {', '.join(wb.sheetnames)}")
